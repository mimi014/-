"""
YouTube文字起こしAPI - Flaskバックエンドサーバー

必要なパッケージ:
- Flask
- flask-cors
- yt-dlp (YouTube動画ダウンロード)
- openai (OpenAI Whisper API)
- google-cloud-speech (オプション: Google Speech-to-Text)
- reportlab (PDF生成)
- python-docx (DOCX生成)
"""

from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import os
import tempfile
import subprocess
import json
from datetime import datetime
from pathlib import Path

app = Flask(__name__)
CORS(app)  # CORS対応

# 設定
UPLOAD_FOLDER = 'uploads'
OUTPUT_FOLDER = 'outputs'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# OpenAI API Key（環境変数から取得）
OPENAI_API_KEY = os.getenv('OPENAI_API_KEY', '')


class YouTubeTranscriber:
    """YouTube動画の文字起こしを行うクラス"""

    def __init__(self, video_url, options=None):
        self.video_url = video_url
        self.options = options or {}
        self.video_id = self._extract_video_id(video_url)
        self.temp_dir = tempfile.mkdtemp()

    def _extract_video_id(self, url):
        """YouTube動画IDを抽出"""
        import re
        regex = r'(?:youtube\.com\/watch\?v=|youtu\.be\/)([^&\n?#]+)'
        match = re.search(regex, url)
        return match.group(1) if match else None

    def download_audio(self):
        """YouTube動画から音声を抽出"""
        audio_path = os.path.join(self.temp_dir, f'{self.video_id}.mp3')

        try:
            # yt-dlpを使用して音声をダウンロード
            cmd = [
                'yt-dlp',
                '-x',  # 音声のみ抽出
                '--audio-format', 'mp3',
                '--audio-quality', '0',
                '-o', audio_path,
                self.video_url
            ]
            subprocess.run(cmd, check=True, capture_output=True)
            return audio_path
        except subprocess.CalledProcessError as e:
            raise Exception(f"音声のダウンロードに失敗しました: {e}")

    def get_video_info(self):
        """動画情報を取得"""
        try:
            cmd = [
                'yt-dlp',
                '--dump-json',
                '--no-download',
                self.video_url
            ]
            result = subprocess.run(cmd, check=True, capture_output=True, text=True)
            info = json.loads(result.stdout)

            return {
                'title': info.get('title', ''),
                'duration': info.get('duration', 0),
                'channel': info.get('uploader', ''),
                'thumbnail': info.get('thumbnail', ''),
                'description': info.get('description', '')
            }
        except Exception as e:
            raise Exception(f"動画情報の取得に失敗しました: {e}")

    def transcribe_audio(self, audio_path):
        """音声を文字起こし（OpenAI Whisper API使用）"""
        try:
            # OpenAI Whisperを使用する場合
            if OPENAI_API_KEY:
                import openai
                openai.api_key = OPENAI_API_KEY

                with open(audio_path, 'rb') as audio_file:
                    transcript = openai.Audio.transcribe(
                        model="whisper-1",
                        file=audio_file,
                        language=self.options.get('sourceLanguage', 'auto'),
                        response_format='verbose_json'
                    )

                return self._process_whisper_response(transcript)
            else:
                # デモモード: ダミーデータを返す
                return self._generate_dummy_transcript()

        except Exception as e:
            raise Exception(f"文字起こしに失敗しました: {e}")

    def _process_whisper_response(self, response):
        """Whisperのレスポンスを処理"""
        segments = []
        full_text = response.get('text', '')

        for segment in response.get('segments', []):
            segments.append({
                'start': segment['start'],
                'end': segment['end'],
                'text': segment['text'].strip(),
                'speaker': None  # 話者識別は別途実装
            })

        return {
            'fullText': full_text,
            'segments': segments
        }

    def _generate_dummy_transcript(self):
        """ダミーの文字起こしデータを生成（デモ用）"""
        return {
            'fullText': 'これはダミーの文字起こしテキストです。実際の動画を文字起こしするには、OpenAI APIキーが必要です。',
            'segments': [
                {'start': 0, 'end': 5, 'text': 'これはダミーの文字起こしテキストです。', 'speaker': 'スピーカー 1'},
                {'start': 5, 'end': 10, 'text': '実際の動画を文字起こしするには、OpenAI APIキーが必要です。', 'speaker': 'スピーカー 1'}
            ]
        }

    def detect_speakers(self, segments):
        """話者識別（簡易版）"""
        # より高度な話者識別には、pyannote.audioなどを使用
        # ここでは簡易的に実装
        for i, segment in enumerate(segments):
            segment['speaker'] = f'スピーカー {(i % 2) + 1}'
        return segments

    def generate_summary(self, text):
        """要約を生成（OpenAI GPT使用）"""
        try:
            if OPENAI_API_KEY:
                import openai
                openai.api_key = OPENAI_API_KEY

                response = openai.ChatCompletion.create(
                    model="gpt-4",
                    messages=[
                        {"role": "system", "content": "あなたは文字起こしテキストを簡潔に要約する専門家です。"},
                        {"role": "user", "content": f"以下のテキストを3-5文で要約してください:\n\n{text}"}
                    ],
                    max_tokens=500
                )

                return response.choices[0].message.content
            else:
                return "要約機能を使用するには、OpenAI APIキーが必要です。"

        except Exception as e:
            raise Exception(f"要約の生成に失敗しました: {e}")

    def translate_text(self, text, target_language):
        """テキストを翻訳（OpenAI GPT使用）"""
        try:
            if OPENAI_API_KEY:
                import openai
                openai.api_key = OPENAI_API_KEY

                lang_names = {
                    'ja': '日本語',
                    'en': '英語',
                    'zh': '中国語',
                    'ko': '韓国語',
                    'es': 'スペイン語',
                    'fr': 'フランス語',
                    'de': 'ドイツ語'
                }

                target_lang_name = lang_names.get(target_language, '英語')

                response = openai.ChatCompletion.create(
                    model="gpt-4",
                    messages=[
                        {"role": "system", "content": f"あなたは優秀な翻訳者です。"},
                        {"role": "user", "content": f"以下のテキストを{target_lang_name}に翻訳してください:\n\n{text}"}
                    ],
                    max_tokens=2000
                )

                return response.choices[0].message.content
            else:
                return "翻訳機能を使用するには、OpenAI APIキーが必要です。"

        except Exception as e:
            raise Exception(f"翻訳に失敗しました: {e}")


# ========== APIエンドポイント ==========

@app.route('/api/transcribe', methods=['POST'])
def transcribe():
    """文字起こしAPIエンドポイント"""
    try:
        data = request.json
        video_url = data.get('videoUrl')
        options = {
            'sourceLanguage': data.get('sourceLanguage', 'auto'),
            'translateTo': data.get('translateTo', 'none'),
            'generateSummary': data.get('generateSummary', False),
            'speakerDetection': data.get('speakerDetection', False)
        }

        if not video_url:
            return jsonify({'error': 'YouTube URLが指定されていません'}), 400

        # 文字起こし処理
        transcriber = YouTubeTranscriber(video_url, options)

        # 動画情報の取得
        video_info = transcriber.get_video_info()

        # 音声のダウンロード
        audio_path = transcriber.download_audio()

        # 文字起こし
        transcript_data = transcriber.transcribe_audio(audio_path)

        # 話者識別
        if options['speakerDetection']:
            transcript_data['segments'] = transcriber.detect_speakers(
                transcript_data['segments']
            )

        # 要約生成
        summary = None
        if options['generateSummary']:
            summary = transcriber.generate_summary(transcript_data['fullText'])

        # 翻訳
        translation = None
        if options['translateTo'] != 'none':
            translation = transcriber.translate_text(
                transcript_data['fullText'],
                options['translateTo']
            )

        # レスポンスの構築
        response_data = {
            'videoInfo': video_info,
            'transcript': transcript_data,
            'summary': summary,
            'translation': translation,
            'timestamp': datetime.now().isoformat()
        }

        return jsonify(response_data), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/export/pdf', methods=['POST'])
def export_pdf():
    """PDF エクスポートエンドポイント"""
    try:
        data = request.json
        text = data.get('text', '')
        title = data.get('title', 'Transcript')

        # PDF生成（reportlabを使用）
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        # 日本語フォントの登録
        pdfmetrics.registerFont(UnicodeCIDFont('HeiseiMin-W3'))

        pdf_path = os.path.join(OUTPUT_FOLDER, f'transcript_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf')

        c = canvas.Canvas(pdf_path, pagesize=A4)
        c.setFont('HeiseiMin-W3', 12)

        # タイトル
        c.drawString(50, 800, title)
        c.drawString(50, 780, '-' * 50)

        # テキストの描画
        y = 750
        for line in text.split('\n'):
            if y < 50:
                c.showPage()
                y = 800
            c.drawString(50, y, line[:80])  # 1行80文字まで
            y -= 20

        c.save()

        return send_file(pdf_path, as_attachment=True)

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/export/docx', methods=['POST'])
def export_docx():
    """DOCX エクスポートエンドポイント"""
    try:
        data = request.json
        text = data.get('text', '')
        title = data.get('title', 'Transcript')

        # DOCX生成（python-docxを使用）
        from docx import Document

        doc = Document()
        doc.add_heading(title, 0)
        doc.add_paragraph(text)

        docx_path = os.path.join(OUTPUT_FOLDER, f'transcript_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx')
        doc.save(docx_path)

        return send_file(docx_path, as_attachment=True)

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/health', methods=['GET'])
def health_check():
    """ヘルスチェックエンドポイント"""
    return jsonify({
        'status': 'ok',
        'timestamp': datetime.now().isoformat(),
        'openai_configured': bool(OPENAI_API_KEY)
    }), 200


# ========== メイン ==========

if __name__ == '__main__':
    print('=' * 50)
    print('YouTube文字起こしAPIサーバー')
    print('=' * 50)
    print(f'OpenAI API設定: {"✓" if OPENAI_API_KEY else "✗"}')
    print('サーバー起動中...')
    print('URL: http://localhost:5000')
    print('=' * 50)

    app.run(debug=True, host='0.0.0.0', port=5000)
